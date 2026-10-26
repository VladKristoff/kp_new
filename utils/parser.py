import os
import re
import sqlite3
from collections import Counter
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config import DB_PATH
from utils.dates import to_iso


class ParseError(Exception):
    """Файл не подходит под ожидаемый формат. Текст ошибки можно показывать пользователю."""


class UniversityParser:
    """
    Парсер Excel-файла с оценками студентов и загрузка их в SQLite.

    Данные пишутся во временную базу и заменяют рабочую только при успехе,
    поэтому неверный файл не уничтожает ранее загруженные данные.

        parser = UniversityParser(excel_file="file.xlsx")
        stats = parser.parse()   # {"students", "activities", "grades", "warnings"}
    """

    IGNORED_FIO_VALUES = {
        "ОТЛИЧНО", "ХОРОШО", "УДОВЛЕТВ.", "НЕУДОВЛ",
        "ИТОГ", "ПРОМЕЖУТОЧНЫЙ БАЛЛ",
    }
    SERVICE_COLUMN_KEYWORDS = ["фио", "фамилия", "итог", "пропуск", "доп.задания"]

    SHEET_DATES = "Даты_темы"
    SHEET_GRADES = "Баллы, оценки"

    # Правила определения типа работы: (проверка по названию, тип, макс. балл по умолчанию)
    ACTIVITY_RULES = [
        (lambda n: re.match(r"^пр\s*\d", n), "ПР", 5),
        (lambda n: re.match(r"^лр\s*\d", n), "ЛР", 5),
        (lambda n: "зачет" in n or "зачёт" in n, "Зачет", 10),
        (lambda n: "оценка" in n, "Оценка", 5),
        (lambda n: "дз" in n, "ДЗ", 3),
    ]
    DEFAULT_TYPE = ("Урок", 2)

    def __init__(self, excel_file: str, db_file: str = DB_PATH, verbose: bool = True):
        self.excel_file = excel_file
        self.db_file = db_file
        self.tmp_file = db_file + ".tmp"
        self.verbose = verbose

        self.conn = None
        self.cursor = None
        self.wb = None

        self.lessons_info = {}       # {номер_урока: {"дата": iso, "тема": str}}
        self.col_to_activity_id = {}  # {индекс_столбца: activity_id}
        self.warnings = []

        self.main_year = None
        self.date_min = None
        self.date_max = None

    # ------------------------------------------------------------------
    # Публичный метод
    # ------------------------------------------------------------------

    def parse(self) -> dict:
        try:
            self._log(f"Старт парсинга: {self.excel_file}")
            self._open_workbook()
            self._check_sheets()
            self._prepare_database()
            self._detect_year_range()
            self._parse_lessons_info()
            self._parse_activities()
            stats = self._parse_students_and_grades()
            self.conn.commit()
            self._close_connections()
            self._replace_database()

            stats["warnings"] = self.warnings
            self._log(f"Студентов: {stats['students']}, работ: {stats['activities']}, "
                      f"оценок: {stats['grades']}, предупреждений: {len(self.warnings)}")
            return stats
        except Exception:
            self._close_connections()
            self._remove_tmp()
            raise
        finally:
            self._close_connections()

    # ------------------------------------------------------------------
    # Подготовка
    # ------------------------------------------------------------------

    def _log(self, message: str):
        if self.verbose:
            print(message)

    def _warn(self, message: str):
        self.warnings.append(message)
        self._log(f"  ! {message}")

    def _open_workbook(self):
        try:
            self.wb = load_workbook(self.excel_file, data_only=True)
        except PermissionError:
            raise ParseError("Файл занят другой программой. Закройте его в Excel и повторите.")
        except Exception as e:
            raise ParseError(f"Не удалось открыть файл как .xlsx: {e}")

    def _check_sheets(self):
        missing = [s for s in (self.SHEET_DATES, self.SHEET_GRADES) if s not in self.wb.sheetnames]
        if missing:
            raise ParseError(
                "Файл не подходит под формат: не найдены листы "
                + ", ".join(f"«{s}»" for s in missing)
                + f".\nНайдены листы: {', '.join(self.wb.sheetnames)}"
            )

    def _prepare_database(self):
        self._remove_tmp()
        self.conn = sqlite3.connect(self.tmp_file)
        self.cursor = self.conn.cursor()
        self.cursor.executescript('''
            CREATE TABLE Students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fio TEXT NOT NULL,
                itog_ball REAL,
                kol_propuskov INTEGER
            );
            CREATE TABLE Activities (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nazvanie TEXT NOT NULL,
                tip TEXT NOT NULL,
                data TEXT,
                tema TEXT,
                max_ball REAL
            );
            CREATE TABLE Grades (
                student_id INTEGER,
                activity_id INTEGER,
                poluchenny_ball REAL,
                raw_value TEXT,
                PRIMARY KEY (student_id, activity_id),
                FOREIGN KEY (student_id) REFERENCES Students(id),
                FOREIGN KEY (activity_id) REFERENCES Activities(id)
            );
            CREATE INDEX idx_grades_activity ON Grades(activity_id);
        ''')

    def _replace_database(self):
        try:
            os.replace(self.tmp_file, self.db_file)
        except PermissionError:
            raise ParseError(
                f"Файл базы данных {os.path.basename(self.db_file)} занят другой программой "
                f"(например, DB Browser for SQLite). Закройте её и повторите."
            )

    def _remove_tmp(self):
        try:
            if os.path.exists(self.tmp_file):
                os.remove(self.tmp_file)
        except OSError:
            pass

    def _close_connections(self):
        if self.conn:
            self.conn.close()
            self.conn = None
        if self.wb:
            self.wb.close()
            self.wb = None

    # ------------------------------------------------------------------
    # Даты
    # ------------------------------------------------------------------

    def _detect_year_range(self):
        """Определяет основной год и диапазон дат семестра по шапке листа с оценками."""
        sheet = self.wb[self.SHEET_GRADES]
        dates = [c.value for c in sheet[2] if isinstance(c.value, datetime)]
        if not dates:
            return
        self.main_year = Counter(d.year for d in dates).most_common(1)[0][0]
        same_year = [d.date() for d in dates if d.year == self.main_year]
        self.date_min, self.date_max = min(same_year), max(same_year)

    def _to_iso(self, value, context: str = "") -> str:
        iso = to_iso(value, default_year=self.main_year)
        return self._fix_year_typo(iso, context)

    def _fix_year_typo(self, iso: str, context: str) -> str:
        """Исправляет опечатки в годе (2026 вместо 2025), если исправленная дата попадает в семестр."""
        if not iso or self.main_year is None or self.date_min is None:
            return iso
        year = int(iso[:4])
        if year == self.main_year:
            return iso
        try:
            fixed = datetime.strptime(f"{self.main_year}{iso[4:]}", "%Y-%m-%d").date()
        except ValueError:
            return iso
        if self.date_min <= fixed <= self.date_max:
            self._warn(f"Дата {iso}{' (' + context + ')' if context else ''} "
                       f"похожа на опечатку в году, заменена на {fixed.isoformat()}")
            return fixed.isoformat()
        return iso

    # ------------------------------------------------------------------
    # Лист «Даты_темы»
    # ------------------------------------------------------------------

    def _parse_lessons_info(self):
        sheet = self.wb[self.SHEET_DATES]
        last_lesson = None

        for row in sheet.iter_rows(min_row=5, values_only=True):
            if len(row) < 4:
                continue
            lesson_raw, date_val, tema_raw = row[1], row[2], row[3]
            tema_part = str(tema_raw).strip() if tema_raw else ""

            lesson_num = None
            if lesson_raw is not None:
                try:
                    lesson_num = int(lesson_raw)
                except (ValueError, TypeError):
                    lesson_num = None

            if lesson_num is not None:
                last_lesson = lesson_num
                info = self.lessons_info.setdefault(
                    lesson_num, {"дата": self._to_iso(date_val, f"урок {lesson_num}"), "тема": ""})
                target = lesson_num
            elif last_lesson is not None:
                target = last_lesson  # продолжение темы в строке без номера урока
            else:
                continue

            if tema_part and tema_part not in self.lessons_info[target]["тема"]:
                current = self.lessons_info[target]["тема"]
                self.lessons_info[target]["тема"] = (current + " " + tema_part).strip()

        self._log(f"Найдено уроков с темами: {len(self.lessons_info)}")

    # ------------------------------------------------------------------
    # Лист «Баллы, оценки» → Activities
    # ------------------------------------------------------------------

    def _parse_activities(self):
        sheet = self.wb[self.SHEET_GRADES]
        header1 = [c.value for c in sheet[1]]
        header2 = [c.value for c in sheet[2]]

        for col_idx, name in enumerate(header1):
            if name is None:
                continue
            name_str = str(name).strip()
            if not name_str:
                continue
            if any(kw in name_str.lower() for kw in self.SERVICE_COLUMN_KEYWORDS):
                continue

            date_iso = self._to_iso(header2[col_idx] if col_idx < len(header2) else None, name_str)
            activity_id = self._insert_activity(name_str, date_iso)
            self.col_to_activity_id[col_idx] = activity_id

        self.conn.commit()
        self._log(f"Видов работ: {len(self.col_to_activity_id)}")

    def _detect_type(self, name: str):
        name_lower = name.lower()
        for check, tip, max_ball in self.ACTIVITY_RULES:
            if check(name_lower):
                return tip, max_ball
        return self.DEFAULT_TYPE

    def _insert_activity(self, name: str, date_iso: str) -> int:
        tip, default_max = self._detect_type(name)

        tema = ""
        match = (re.search(r"(?:ур\.?|урок)\s*(\d+)", name, re.IGNORECASE)
                 or re.search(r"^(\d+)$", name))
        if match:
            lesson_num = int(match.group(1))
            if lesson_num in self.lessons_info:
                tema = self.lessons_info[lesson_num]["тема"]
                if not date_iso:
                    date_iso = self.lessons_info[lesson_num]["дата"]

        # Для итоговых столбцов («оценка 1», зачёт) тему по дате не подбираем — она будет чужой
        if not tema and date_iso and tip not in ("Оценка", "Зачет"):
            for info in self.lessons_info.values():
                if info["дата"] == date_iso:
                    tema = info["тема"]
                    break

        self.cursor.execute(
            "INSERT INTO Activities (nazvanie, tip, data, tema, max_ball) VALUES (?, ?, ?, ?, ?)",
            (name, tip, date_iso, tema, default_max))
        return self.cursor.lastrowid

    # ------------------------------------------------------------------
    # Студенты и оценки
    # ------------------------------------------------------------------

    def _parse_students_and_grades(self) -> dict:
        sheet = self.wb[self.SHEET_GRADES]
        header1 = [c.value for c in sheet[1]]
        fio_col, itog_col, propusk_col = self._find_service_columns(header1)

        if fio_col is None:
            raise ParseError("В листе «Баллы, оценки» не найден столбец с ФИО "
                             "(ожидается заголовок «ФИО» или «Фамилия ...»).")
        self._log(f"Столбцы: ФИО={get_column_letter(fio_col + 1)}, "
                  f"итог={get_column_letter(itog_col + 1) if itog_col is not None else '—'}, "
                  f"пропуски={get_column_letter(propusk_col + 1) if propusk_col is not None else '—'}")
        if itog_col is None:
            self._warn("Не найден столбец «итог» — итоговые баллы будут пустыми")
        if propusk_col is None:
            self._warn("Не найден столбец «пропуски» — число пропусков будет пустым")

        students = grades_total = 0
        unrecognized = Counter()
        observed_max = {}

        for row in sheet.iter_rows(min_row=3, values_only=True):
            student_id = self._insert_student(row, fio_col, itog_col, propusk_col)
            if student_id is None:
                continue
            students += 1
            grades_total += self._insert_grades(student_id, row, unrecognized, observed_max)

        if students == 0:
            raise ParseError("В листе «Баллы, оценки» не найдено ни одного студента.")

        # Максимальный балл — фактический максимум среди студентов (точнее догадки по названию)
        for col_idx, best in observed_max.items():
            self.cursor.execute("UPDATE Activities SET max_ball = ? WHERE id = ?",
                                (best, self.col_to_activity_id[col_idx]))

        for text, count in unrecognized.most_common(5):
            self._warn(f"Нечисловое значение «{text}» в оценках ({count} шт.) — "
                       f"сохранено как есть, пропуском не считается")
        return {"students": students, "activities": len(self.col_to_activity_id),
                "grades": grades_total}

    def _find_service_columns(self, header1):
        fio_col = itog_col = propusk_col = None
        for idx, val in enumerate(header1):
            if val is None:
                continue
            text = str(val).strip().lower()
            if fio_col is None and ("фио" in text or "фамилия" in text):
                fio_col = idx
            elif itog_col is None and "итог" in text:
                itog_col = idx
            elif propusk_col is None and "пропуск" in text:
                propusk_col = idx
        return fio_col, itog_col, propusk_col

    def _insert_student(self, row, fio_col, itog_col, propusk_col):
        if len(row) <= fio_col or row[fio_col] is None:
            return None
        fio = " ".join(str(row[fio_col]).split())
        if not fio or fio.upper() in self.IGNORED_FIO_VALUES:
            return None

        itog = self._parse_float(row, itog_col)
        propuski = self._parse_int(row, propusk_col)
        if itog is None and propuski is None:
            return None  # служебная строка (легенда, комментарии)

        self.cursor.execute(
            "INSERT INTO Students (fio, itog_ball, kol_propuskov) VALUES (?, ?, ?)",
            (fio, itog, propuski))
        return self.cursor.lastrowid

    def _insert_grades(self, student_id: int, row, unrecognized: Counter, observed_max: dict) -> int:
        batch = []
        for col_idx, activity_id in self.col_to_activity_id.items():
            if col_idx >= len(row):
                continue
            val = row[col_idx]
            ball = self._parse_float(row, col_idx)
            raw = None
            if ball is None and val is not None and str(val).strip() != "":
                raw = str(val).strip()   # нечисловое значение: «5-», «#NAME?», формула
                unrecognized[raw] += 1
            if ball is not None and ball > observed_max.get(col_idx, float("-inf")):
                observed_max[col_idx] = ball
            batch.append((student_id, activity_id, ball, raw))
        self.cursor.executemany(
            "INSERT INTO Grades (student_id, activity_id, poluchenny_ball, raw_value) "
            "VALUES (?, ?, ?, ?)", batch)
        return len(batch)

    # ------------------------------------------------------------------
    # Утилиты
    # ------------------------------------------------------------------

    @staticmethod
    def _parse_float(row, col_idx):
        if col_idx is None or col_idx >= len(row):
            return None
        val = row[col_idx]
        if val is None or isinstance(val, bool):
            return None
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str):
            try:
                return float(val.strip().replace(",", "."))
            except ValueError:
                return None
        return None

    @staticmethod
    def _parse_int(row, col_idx):
        value = UniversityParser._parse_float(row, col_idx)
        return int(value) if value is not None else None
